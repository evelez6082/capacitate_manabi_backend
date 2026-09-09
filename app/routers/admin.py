from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any
import unicodedata

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from psycopg import Connection

from app.db import fetch_all, fetch_one, get_connection
from app.security import require_roles

router = APIRouter(prefix="/api/admin", tags=["admin"])

CANTON_COORDS = {
    "24 de mayo": (-1.2798, -80.4182),
    "bolivar": (-0.8442, -80.1649),
    "chone": (-0.6982, -80.0936),
    "el carmen": (-0.2689, -79.4596),
    "flavio alfaro": (-0.4048, -79.9034),
    "jama": (-0.2048, -80.2647),
    "jaramijo": (-0.9472, -80.6408),
    "jipijapa": (-1.3487, -80.5788),
    "junin": (-0.9272, -80.2058),
    "manta": (-0.9677, -80.7089),
    "montecristi": (-1.0458, -80.6589),
    "olmedo": (-1.3671, -80.2074),
    "pajan": (-1.5524, -80.4286),
    "pedernales": (0.0719, -80.0528),
    "pichincha": (-1.0499, -79.8175),
    "portoviejo": (-1.0546, -80.4545),
    "puerto lopez": (-1.5604, -80.8116),
    "rocafuerte": (-0.9236, -80.4495),
    "san vicente": (-0.6049, -80.4023),
    "santa ana": (-1.2075, -80.3712),
    "sucre": (-0.5981, -80.4248),
    "tosagua": (-0.7867, -80.2347),
}
MANABI_CENTER = (-1.0546, -80.4545)

CONSULTA_COLUMNS = [
    "cedula",
    "nombre_completo",
    "correo_principal",
    "telefono_principal",
    "provincia",
    "canton",
    "parroquia",
    "curso",
    "version_moodle",
    "campana_inscripcion",
    "fecha_inscripcion",
    "porcentaje_avance",
    "estado_avance",
    "estado_aprobacion",
    "cohorte_aprobacion",
    "fecha_aprobacion",
    "estado_solicitud_diploma",
    "fecha_solicitud_diploma",
    "numero_diploma",
    "diploma_url",
]

CONSULTA_LABELS = {
    "inscritos": "Inscritos",
    "aprobados": "Aprobados",
    "con-diploma": "Con diploma",
    "en-curso": "En curso sin aprobar",
    "aprobados-sin-diploma": "Aprobados sin diploma",
}

CONSULTA_WHERE = {
    "inscritos": "i.id IS NOT NULL",
    "aprobados": "a.id IS NOT NULL",
    "con-diploma": "d.id IS NOT NULL",
    "en-curso": "i.id IS NOT NULL AND a.id IS NULL AND d.id IS NULL",
    "aprobados-sin-diploma": "a.id IS NOT NULL AND d.id IS NULL",
}

INSCRIPCIONES_EXPORT_COLUMNS = [
    "inscripcion_id",
    "persona_id",
    "cedula",
    "nombres",
    "apellidos",
    "nombre_completo",
    "correo_principal",
    "telefono_principal",
    "fecha_nacimiento",
    "genero",
    "etnia",
    "nivel_educativo",
    "discapacidad",
    "nacionalidad",
    "provincia",
    "canton",
    "parroquia",
    "sector",
    "curso",
    "version_moodle",
    "campana_inscripcion",
    "campana_codigo",
    "campana_slug",
    "fecha_inscripcion",
    "modalidad",
    "ocupacion",
    "institucion",
    "estado",
    "consentimiento",
    "observacion",
    "source_file_id",
    "import_batch_id",
    "submission_id",
    "raw_data",
    "created_at",
]

MOODLE_EXPORT_COLUMNS = [
    "username",
    "password",
    "firstname",
    "lastname",
    "email",
    "phone1",
    "country",
    "timezone",
    "lang",
    "cohort1",
]


def normalize_key(value: str | None) -> str:
    if not value:
        return ""
    normalized = unicodedata.normalize("NFD", value)
    without_accents = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    return without_accents.lower().strip()


def consulta_base_sql(where_clause: str, has_search: bool = False) -> str:
    search_clause = ""
    if has_search:
        search_clause = """
          AND (
              p.cedula ILIKE %(search)s
              OR p.nombre_completo ILIKE %(search)s
              OR p.correo_principal ILIKE %(search)s
              OR p.telefono_principal ILIKE %(search)s
          )
        """

    return f"""
        WITH persona_cursos AS (
            SELECT persona_id, curso_id FROM inscripciones
            UNION
            SELECT persona_id, curso_id FROM aprobaciones
            UNION
            SELECT persona_id, curso_id FROM diplomas
        ),
        latest_inscripciones AS (
            SELECT DISTINCT ON (persona_id, curso_id) *
            FROM inscripciones
            ORDER BY persona_id, curso_id, fecha_inscripcion DESC NULLS LAST, id DESC
        )
        SELECT
            p.cedula,
            p.nombre_completo,
            p.correo_principal,
            p.telefono_principal,
            pr.nombre AS provincia,
            ct.nombre AS canton,
            pa.nombre AS parroquia,
            c.nombre AS curso,
            cv.nombre AS version_moodle,
            ci.nombre AS campana_inscripcion,
            i.fecha_inscripcion,
            aar.porcentaje_avance,
            aar.estado_general AS estado_avance,
            a.estado AS estado_aprobacion,
            ca.etiqueta AS cohorte_aprobacion,
            a.fecha_aprobacion,
            sd.estado AS estado_solicitud_diploma,
            lsd.fecha_solicitud AS fecha_solicitud_diploma,
            d.numero_diploma,
            d.archivo_url AS diploma_url
        FROM persona_cursos pc
        JOIN personas p ON p.id = pc.persona_id
        JOIN cursos c ON c.id = pc.curso_id
        LEFT JOIN latest_inscripciones i ON i.persona_id = pc.persona_id AND i.curso_id = pc.curso_id
        LEFT JOIN provincias pr ON pr.id = p.provincia_id
        LEFT JOIN cantones ct ON ct.id = p.canton_id
        LEFT JOIN parroquias pa ON pa.id = p.parroquia_id
        LEFT JOIN curso_versiones_moodle cv ON cv.id = i.curso_version_id
        LEFT JOIN campanas_inscripcion ci ON ci.id = i.campana_inscripcion_id
        LEFT JOIN LATERAL (
            SELECT aer.*
            FROM avance_estudiante_resumen aer
            JOIN campanas_avance cav ON cav.id = aer.campana_avance_id
            WHERE aer.persona_id = p.id
              AND (i.curso_version_id IS NULL OR cav.curso_version_id = i.curso_version_id)
            ORDER BY cav.fecha_corte DESC NULLS LAST, aer.id DESC
            LIMIT 1
        ) aar ON true
        LEFT JOIN LATERAL (
            SELECT ap.*
            FROM aprobaciones ap
            WHERE ap.persona_id = p.id
              AND ap.curso_id = pc.curso_id
              AND ap.estado = 'aprobado'
            ORDER BY ap.fecha_aprobacion DESC NULLS LAST, ap.created_at DESC
            LIMIT 1
        ) a ON true
        LEFT JOIN cohortes_aprobacion ca ON ca.id = a.cohorte_aprobacion_id
        LEFT JOIN LATERAL (
            SELECT sol.*
            FROM solicitudes_diploma sol
            WHERE sol.persona_id = p.id
              AND sol.curso_id = pc.curso_id
              AND sol.estado <> 'anulado'
            ORDER BY sol.created_at DESC
            LIMIT 1
        ) sd ON true
        LEFT JOIN lotes_solicitud_diploma lsd ON lsd.id = sd.lote_solicitud_id
        LEFT JOIN diplomas d ON d.persona_id = p.id AND d.curso_id = pc.curso_id AND d.estado <> 'anulado'
        WHERE {where_clause}
        {search_clause}
    """


def consulta_params(q: str | None) -> dict[str, Any]:
    return {"search": f"%{q.strip()}%"} if q and q.strip() else {}


def serialize_csv_value(value: Any) -> Any:
    if value is None:
        return ""
    return value


def rows_to_csv(rows: list[dict[str, Any]]) -> str:
    output = StringIO()
    writer = csv.DictWriter(output, fieldnames=CONSULTA_COLUMNS, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: serialize_csv_value(row.get(key)) for key in CONSULTA_COLUMNS})
    return output.getvalue()


def serialize_excel_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        import json

        return json.dumps(value, ensure_ascii=False)
    return value


def rows_to_xlsx(rows: list[dict[str, Any]], columns: list[str], sheet_name: str) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(columns)

    header_fill = PatternFill("solid", fgColor="D9EDE9")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        worksheet.append([serialize_excel_value(row.get(column)) for column in columns])

    for index, column in enumerate(columns, start=1):
        max_length = max(len(str(column)), *(len(str(serialize_excel_value(row.get(column)))) for row in rows[:200]))
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 42)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def xlsx_response(stream: BytesIO, filename: str) -> StreamingResponse:
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def inscripciones_base_where(
    anio: int | None,
    mes: int | None,
    campana_id: int | None,
    q: str | None,
) -> tuple[str, dict[str, Any]]:
    clauses = ["TRUE"]
    params: dict[str, Any] = {}
    if anio:
        clauses.append("EXTRACT(YEAR FROM i.fecha_inscripcion)::int = %(anio)s")
        params["anio"] = anio
    if mes:
        clauses.append("EXTRACT(MONTH FROM i.fecha_inscripcion)::int = %(mes)s")
        params["mes"] = mes
    if campana_id:
        clauses.append("i.campana_inscripcion_id = %(campana_id)s")
        params["campana_id"] = campana_id
    if q and q.strip():
        clauses.append(
            """
            (
                p.cedula ILIKE %(search)s
                OR p.nombre_completo ILIKE %(search)s
                OR p.correo_principal ILIKE %(search)s
                OR p.telefono_principal ILIKE %(search)s
            )
            """
        )
        params["search"] = f"%{q.strip()}%"
    return " AND ".join(clauses), params


def inscripciones_select_sql(where_clause: str) -> str:
    return f"""
        SELECT
            i.id AS inscripcion_id,
            p.id AS persona_id,
            p.cedula,
            p.nombre_completo,
            p.correo_principal,
            p.telefono_principal,
            pr.nombre AS provincia,
            ct.nombre AS canton,
            pa.nombre AS parroquia,
            c.nombre AS curso,
            cv.nombre AS version_moodle,
            ci.id AS campana_inscripcion_id,
            ci.nombre AS campana_inscripcion,
            i.fecha_inscripcion,
            i.estado,
            i.modalidad,
            i.ocupacion,
            i.institucion
        FROM inscripciones i
        JOIN personas p ON p.id = i.persona_id
        JOIN cursos c ON c.id = i.curso_id
        LEFT JOIN provincias pr ON pr.id = p.provincia_id
        LEFT JOIN cantones ct ON ct.id = p.canton_id
        LEFT JOIN parroquias pa ON pa.id = p.parroquia_id
        LEFT JOIN curso_versiones_moodle cv ON cv.id = i.curso_version_id
        LEFT JOIN campanas_inscripcion ci ON ci.id = i.campana_inscripcion_id
        WHERE {where_clause}
    """


def inscripciones_export_sql(where_clause: str) -> str:
    return f"""
        SELECT
            i.id AS inscripcion_id,
            p.id AS persona_id,
            p.cedula,
            p.nombres,
            p.apellidos,
            p.nombre_completo,
            p.correo_principal,
            p.telefono_principal,
            p.fecha_nacimiento,
            p.genero,
            p.etnia,
            p.nivel_educativo,
            p.discapacidad,
            n.nombre AS nacionalidad,
            pr.nombre AS provincia,
            ct.nombre AS canton,
            pa.nombre AS parroquia,
            p.sector,
            c.nombre AS curso,
            cv.nombre AS version_moodle,
            ci.nombre AS campana_inscripcion,
            ci.codigo AS campana_codigo,
            ci.slug_publico AS campana_slug,
            i.fecha_inscripcion,
            i.modalidad,
            i.ocupacion,
            i.institucion,
            i.estado,
            i.consentimiento,
            i.observacion,
            i.source_file_id,
            i.import_batch_id,
            i.submission_id,
            i.raw_data,
            i.created_at
        FROM inscripciones i
        JOIN personas p ON p.id = i.persona_id
        JOIN cursos c ON c.id = i.curso_id
        LEFT JOIN nacionalidades n ON n.id = p.nacionalidad_id
        LEFT JOIN provincias pr ON pr.id = p.provincia_id
        LEFT JOIN cantones ct ON ct.id = p.canton_id
        LEFT JOIN parroquias pa ON pa.id = p.parroquia_id
        LEFT JOIN curso_versiones_moodle cv ON cv.id = i.curso_version_id
        LEFT JOIN campanas_inscripcion ci ON ci.id = i.campana_inscripcion_id
        WHERE {where_clause}
    """


def inscripciones_moodle_sql(where_clause: str) -> str:
    return f"""
        SELECT
            p.cedula AS username,
            p.cedula AS password,
            COALESCE(NULLIF(p.nombres, ''), split_part(p.nombre_completo, ' ', 1)) AS firstname,
            COALESCE(NULLIF(p.apellidos, ''), trim(regexp_replace(p.nombre_completo, '^\\S+\\s*', ''))) AS lastname,
            p.correo_principal AS email,
            p.telefono_principal AS phone1,
            'EC' AS country,
            'America/Guayaquil' AS timezone,
            'es' AS lang,
            COALESCE(cm.codigo, ci.codigo, cv.codigo, ci.slug_publico) AS cohort1
        FROM inscripciones i
        JOIN personas p ON p.id = i.persona_id
        LEFT JOIN curso_versiones_moodle cv ON cv.id = i.curso_version_id
        LEFT JOIN campanas_inscripcion ci ON ci.id = i.campana_inscripcion_id
        LEFT JOIN LATERAL (
            SELECT codigo
            FROM cohortes_matriculacion cm
            WHERE (cm.campana_inscripcion_id = i.campana_inscripcion_id)
               OR (
                    cm.campana_inscripcion_id IS NULL
                    AND cm.curso_version_id = i.curso_version_id
                  )
            ORDER BY
                CASE WHEN cm.campana_inscripcion_id = i.campana_inscripcion_id THEN 0 ELSE 1 END,
                cm.fecha_creacion DESC
            LIMIT 1
        ) cm ON true
        WHERE {where_clause}
    """


def avance_rango_case() -> str:
    return """
        CASE
            WHEN COALESCE(aar.porcentaje_avance, 0) = 0 THEN 'Sin avance'
            WHEN aar.porcentaje_avance < 50 THEN '1%% - 49%%'
            WHEN aar.porcentaje_avance < 100 THEN '50%% - 99%%'
            ELSE '100%%'
        END
    """


def count_consulta(conn: Connection, tipo: str, params: dict[str, Any]) -> int:
    if params:
        where_clause = CONSULTA_WHERE[tipo]
        base_sql = consulta_base_sql(where_clause, has_search=True)
        return fetch_one(conn, f"SELECT count(*) AS total FROM ({base_sql}) t", params)["total"]

    queries = {
        "inscritos": """
            SELECT count(*) AS total
            FROM (
                SELECT DISTINCT persona_id, curso_id
                FROM inscripciones
            ) t
        """,
        "aprobados": """
            SELECT count(*) AS total
            FROM (
                SELECT DISTINCT persona_id, curso_id
                FROM aprobaciones
                WHERE estado = 'aprobado'
            ) t
        """,
        "con-diploma": """
            SELECT count(*) AS total
            FROM (
                SELECT DISTINCT persona_id, curso_id
                FROM diplomas
                WHERE estado <> 'anulado'
            ) t
        """,
        "en-curso": """
            SELECT count(*) AS total
            FROM (
                SELECT DISTINCT i.persona_id, i.curso_id
                FROM inscripciones i
                LEFT JOIN aprobaciones a
                    ON a.persona_id = i.persona_id
                   AND a.curso_id = i.curso_id
                   AND a.estado = 'aprobado'
                LEFT JOIN diplomas d
                    ON d.persona_id = i.persona_id
                   AND d.curso_id = i.curso_id
                   AND d.estado <> 'anulado'
                WHERE a.id IS NULL AND d.id IS NULL
            ) t
        """,
        "aprobados-sin-diploma": """
            SELECT count(*) AS total
            FROM (
                SELECT DISTINCT a.persona_id, a.curso_id
                FROM aprobaciones a
                LEFT JOIN diplomas d
                    ON d.persona_id = a.persona_id
                   AND d.curso_id = a.curso_id
                   AND d.estado <> 'anulado'
                WHERE a.estado = 'aprobado'
                  AND d.id IS NULL
            ) t
        """,
    }
    return fetch_one(conn, queries[tipo])["total"]


@router.get("/metricas/resumen")
def resumen_admin(
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    return {
        "personas": fetch_one(conn, "SELECT count(*) AS total FROM personas")["total"],
        "inscripciones": fetch_one(conn, "SELECT count(*) AS total FROM inscripciones")["total"],
        "campanas_activas": fetch_one(
            conn,
            "SELECT count(*) AS total FROM campanas_inscripcion WHERE estado = 'activa'",
        )["total"],
        "territorios": fetch_one(
            conn,
            """
            SELECT count(*) AS total
            FROM (
                SELECT DISTINCT p.canton_id
                FROM inscripciones i
                JOIN personas p ON p.id = i.persona_id
                WHERE p.canton_id IS NOT NULL
            ) t
            """,
        )["total"],
    }


@router.get("/metricas/inscripciones-territorio")
def inscripciones_territorio(
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    rows = fetch_all(
        conn,
        """
        SELECT
            pr.nombre AS provincia,
            c.nombre AS canton,
            p.parroquia_id,
            pa.nombre AS parroquia,
            count(i.id)::int AS total
        FROM inscripciones i
        JOIN personas p ON p.id = i.persona_id
        LEFT JOIN provincias pr ON pr.id = p.provincia_id
        LEFT JOIN cantones c ON c.id = p.canton_id
        LEFT JOIN parroquias pa ON pa.id = p.parroquia_id
        GROUP BY pr.nombre, c.nombre, p.parroquia_id, pa.nombre
        ORDER BY total DESC, provincia, canton, parroquia
        """,
    )

    by_canton: dict[str, dict[str, Any]] = {}
    for row in rows:
        canton = row["canton"] or "Sin canton"
        key = f"{row['provincia'] or 'Sin provincia'}::{canton}"
        coords = CANTON_COORDS.get(normalize_key(canton), MANABI_CENTER)
        item = by_canton.setdefault(
            key,
            {
                "provincia": row["provincia"] or "Sin provincia",
                "canton": canton,
                "lat": coords[0],
                "lng": coords[1],
                "total": 0,
                "parroquias": [],
            },
        )
        item["total"] += int(row["total"])
        item["parroquias"].append(
            {
                "nombre": row["parroquia"] or "Sin parroquia",
                "total": int(row["total"]),
            }
        )

    items = sorted(by_canton.values(), key=lambda item: item["total"], reverse=True)
    return {
        "items": items,
        "total": sum(item["total"] for item in items),
        "center": {"lat": MANABI_CENTER[0], "lng": MANABI_CENTER[1]},
    }


@router.get("/inscritos/resumen")
def resumen_inscritos(
    agrupar_por: str = Query(default="anio", pattern="^(anio|mes|cohorte)$"),
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    if agrupar_por == "anio":
        rows = fetch_all(
            conn,
            """
            SELECT
                EXTRACT(YEAR FROM fecha_inscripcion)::int AS anio,
                EXTRACT(YEAR FROM fecha_inscripcion)::int AS clave,
                COALESCE(EXTRACT(YEAR FROM fecha_inscripcion)::int::text, 'Sin fecha') AS nombre,
                count(*)::int AS total
            FROM inscripciones
            GROUP BY EXTRACT(YEAR FROM fecha_inscripcion)::int
            ORDER BY anio DESC NULLS LAST
            """,
        )
    elif agrupar_por == "mes":
        rows = fetch_all(
            conn,
            """
            SELECT
                EXTRACT(YEAR FROM fecha_inscripcion)::int AS anio,
                EXTRACT(MONTH FROM fecha_inscripcion)::int AS mes,
                to_char(date_trunc('month', fecha_inscripcion), 'YYYY-MM') AS clave,
                COALESCE(to_char(date_trunc('month', fecha_inscripcion), 'TMMonth YYYY'), 'Sin fecha') AS nombre,
                count(*)::int AS total
            FROM inscripciones
            GROUP BY date_trunc('month', fecha_inscripcion), EXTRACT(YEAR FROM fecha_inscripcion)::int, EXTRACT(MONTH FROM fecha_inscripcion)::int
            ORDER BY date_trunc('month', fecha_inscripcion) DESC NULLS LAST
            """,
        )
    else:
        rows = fetch_all(
            conn,
            """
            SELECT
                ci.id AS campana_inscripcion_id,
                COALESCE(ci.nombre, 'Sin cohorte/campana') AS nombre,
                count(i.id)::int AS total,
                min(i.fecha_inscripcion) AS primera_inscripcion,
                max(i.fecha_inscripcion) AS ultima_inscripcion
            FROM inscripciones i
            LEFT JOIN campanas_inscripcion ci ON ci.id = i.campana_inscripcion_id
            GROUP BY ci.id, ci.nombre
            ORDER BY ultima_inscripcion DESC NULLS LAST, total DESC
            """,
        )
    return {"agrupar_por": agrupar_por, "items": rows}


@router.get("/inscritos")
def listar_inscritos(
    anio: int | None = Query(default=None, ge=2000, le=2100),
    mes: int | None = Query(default=None, ge=1, le=12),
    campana_id: int | None = Query(default=None, ge=1),
    q: str | None = Query(default=None),
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    where_clause, params = inscripciones_base_where(anio, mes, campana_id, q)
    base_sql = inscripciones_select_sql(where_clause)
    total = fetch_one(conn, f"SELECT count(*) AS total FROM ({base_sql}) t", params)["total"]
    rows = fetch_all(
        conn,
        f"""
        SELECT *
        FROM ({base_sql}) t
        ORDER BY fecha_inscripcion DESC NULLS LAST, nombre_completo
        LIMIT %(limit)s OFFSET %(offset)s
        """,
        {**params, "limit": limit, "offset": offset},
    )
    return {"items": rows, "total": total, "limit": limit, "offset": offset}


@router.get("/inscritos/export.xlsx")
def exportar_inscritos_excel(
    anio: int | None = Query(default=None, ge=2000, le=2100),
    mes: int | None = Query(default=None, ge=1, le=12),
    campana_id: int | None = Query(default=None, ge=1),
    q: str | None = Query(default=None),
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> StreamingResponse:
    where_clause, params = inscripciones_base_where(anio, mes, campana_id, q)
    base_sql = inscripciones_export_sql(where_clause)
    rows = fetch_all(
        conn,
        f"""
        SELECT *
        FROM ({base_sql}) t
        ORDER BY fecha_inscripcion DESC NULLS LAST, nombre_completo
        """,
        params,
    )
    stream = rows_to_xlsx(rows, INSCRIPCIONES_EXPORT_COLUMNS, "inscripciones")
    return xlsx_response(stream, "capacitate-manabi-inscripciones.xlsx")


@router.get("/inscritos/export-moodle.xlsx")
def exportar_inscritos_moodle_excel(
    anio: int | None = Query(default=None, ge=2000, le=2100),
    mes: int | None = Query(default=None, ge=1, le=12),
    campana_id: int | None = Query(default=None, ge=1),
    q: str | None = Query(default=None),
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> StreamingResponse:
    where_clause, params = inscripciones_base_where(anio, mes, campana_id, q)
    base_sql = inscripciones_moodle_sql(where_clause)
    rows = fetch_all(
        conn,
        f"""
        SELECT *
        FROM ({base_sql}) t
        ORDER BY username
        """,
        params,
    )
    stream = rows_to_xlsx(rows, MOODLE_EXPORT_COLUMNS, "moodle_users")
    return xlsx_response(stream, "capacitate-manabi-moodle-users.xlsx")


@router.get("/inscritos/{persona_id}/detalle")
def detalle_inscrito(
    persona_id: int,
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    persona = fetch_one(
        conn,
        """
        SELECT
            p.id,
            p.cedula,
            p.nombres,
            p.apellidos,
            p.nombre_completo,
            p.correo_principal,
            p.telefono_principal,
            p.fecha_nacimiento,
            p.genero,
            p.etnia,
            p.nivel_educativo,
            p.discapacidad,
            n.nombre AS nacionalidad,
            pr.nombre AS provincia,
            ct.nombre AS canton,
            pa.nombre AS parroquia,
            p.sector,
            p.created_at
        FROM personas p
        LEFT JOIN nacionalidades n ON n.id = p.nacionalidad_id
        LEFT JOIN provincias pr ON pr.id = p.provincia_id
        LEFT JOIN cantones ct ON ct.id = p.canton_id
        LEFT JOIN parroquias pa ON pa.id = p.parroquia_id
        WHERE p.id = %s
        """,
        (persona_id,),
    )
    if not persona:
        raise HTTPException(status_code=404, detail="Inscrito no encontrado")

    inscripciones = fetch_all(
        conn,
        """
        SELECT
            i.id AS inscripcion_id,
            c.nombre AS curso,
            cv.nombre AS version_moodle,
            ci.nombre AS campana_inscripcion,
            ci.codigo AS campana_codigo,
            i.fecha_inscripcion,
            i.estado,
            i.modalidad,
            i.ocupacion,
            i.institucion,
            i.consentimiento,
            i.observacion,
            i.source_file_id,
            i.import_batch_id,
            i.submission_id,
            i.raw_data,
            i.created_at
        FROM inscripciones i
        JOIN cursos c ON c.id = i.curso_id
        LEFT JOIN curso_versiones_moodle cv ON cv.id = i.curso_version_id
        LEFT JOIN campanas_inscripcion ci ON ci.id = i.campana_inscripcion_id
        WHERE i.persona_id = %s
        ORDER BY i.fecha_inscripcion DESC NULLS LAST, i.id DESC
        """,
        (persona_id,),
    )

    trazabilidad = fetch_all(
        conn,
        """
        SELECT *
        FROM vw_trazabilidad_participante
        WHERE cedula = %s
        ORDER BY fecha_inscripcion DESC NULLS LAST, fecha_solicitud_diploma DESC NULLS LAST
        """,
        (persona["cedula"],),
    )
    return {"persona": persona, "inscripciones": inscripciones, "trazabilidad": trazabilidad}


@router.get("/aprobados-avance/resumen")
def resumen_aprobados_avance(
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    aprobados_por_cohorte = fetch_all(
        conn,
        """
        SELECT
            COALESCE(ca.etiqueta, 'Sin cohorte') AS cohorte,
            ca.anio,
            ca.mes,
            count(DISTINCT (a.persona_id, a.curso_id))::int AS total
        FROM aprobaciones a
        LEFT JOIN cohortes_aprobacion ca ON ca.id = a.cohorte_aprobacion_id
        WHERE a.estado = 'aprobado'
        GROUP BY ca.etiqueta, ca.anio, ca.mes, ca.numero
        ORDER BY ca.anio DESC NULLS LAST, ca.numero DESC NULLS LAST, total DESC
        """,
    )
    aprobados_por_anio = fetch_all(
        conn,
        """
        SELECT
            EXTRACT(YEAR FROM fecha_aprobacion)::int AS anio,
            COALESCE(EXTRACT(YEAR FROM fecha_aprobacion)::int::text, 'Sin fecha') AS nombre,
            count(DISTINCT (persona_id, curso_id))::int AS total
        FROM aprobaciones
        WHERE estado = 'aprobado'
        GROUP BY EXTRACT(YEAR FROM fecha_aprobacion)::int
        ORDER BY anio DESC NULLS LAST
        """,
    )
    avance_por_rango = fetch_all(
        conn,
        f"""
        WITH ultimo_avance AS (
            SELECT DISTINCT ON (persona_id)
                persona_id,
                porcentaje_avance,
                estado_general
            FROM avance_estudiante_resumen
            WHERE persona_id IS NOT NULL
            ORDER BY persona_id, id DESC
        )
        SELECT rango, count(*)::int AS total
        FROM (
            SELECT {avance_rango_case()} AS rango
            FROM inscripciones i
            JOIN personas p ON p.id = i.persona_id
            LEFT JOIN ultimo_avance aar ON aar.persona_id = p.id
            LEFT JOIN aprobaciones a ON a.persona_id = p.id AND a.curso_id = i.curso_id AND a.estado = 'aprobado'
            WHERE a.id IS NULL
        ) t
        GROUP BY rango
        ORDER BY
            CASE rango
                WHEN 'Sin avance' THEN 1
                WHEN '1%% - 49%%' THEN 2
                WHEN '50%% - 99%%' THEN 3
                ELSE 4
            END
        """,
    )
    return {
        "aprobados_por_cohorte": aprobados_por_cohorte,
        "aprobados_por_anio": aprobados_por_anio,
        "avance_por_rango": avance_por_rango,
    }


@router.get("/consultas/resumen")
def resumen_consultas(
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    counts = {key: count_consulta(conn, key, {}) for key in CONSULTA_LABELS}
    return {
        "items": [
            {"codigo": key, "nombre": CONSULTA_LABELS[key], "total": total}
            for key, total in counts.items()
        ]
    }


@router.get("/consultas/{tipo}")
def listar_consulta(
    tipo: str,
    q: str | None = Query(default=None),
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> dict[str, Any]:
    where_clause = CONSULTA_WHERE.get(tipo)
    if where_clause is None:
        raise HTTPException(status_code=404, detail="Consulta no encontrada")

    params = consulta_params(q)
    base_sql = consulta_base_sql(where_clause, has_search=bool(params))
    total = count_consulta(conn, tipo, params)
    rows = fetch_all(
        conn,
        f"""
        SELECT *
        FROM ({base_sql}) t
        ORDER BY fecha_inscripcion DESC NULLS LAST, nombre_completo
        LIMIT %(limit)s OFFSET %(offset)s
        """,
        {**params, "limit": limit, "offset": offset},
    )
    return {
        "tipo": tipo,
        "nombre": CONSULTA_LABELS[tipo],
        "items": rows,
        "total": total,
        "limit": limit,
        "offset": offset,
    }


@router.get("/consultas/{tipo}/export.csv")
def exportar_consulta_csv(
    tipo: str,
    q: str | None = Query(default=None),
    conn: Connection = Depends(get_connection),
    _: dict[str, Any] = Depends(require_roles("admin", "supervisor")),
) -> Response:
    where_clause = CONSULTA_WHERE.get(tipo)
    if where_clause is None:
        raise HTTPException(status_code=404, detail="Consulta no encontrada")

    params = consulta_params(q)
    base_sql = consulta_base_sql(where_clause, has_search=bool(params))
    rows = fetch_all(
        conn,
        f"""
        SELECT *
        FROM ({base_sql}) t
        ORDER BY fecha_inscripcion DESC NULLS LAST, nombre_completo
        """,
        params,
    )
    filename = f"capacitate-manabi-{tipo}.csv"
    return Response(
        content=rows_to_csv(rows),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
